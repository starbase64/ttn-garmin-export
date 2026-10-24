import csv
import inspect
import json
import threading
from pathlib import Path
import pytest
from garminconnect import Garmin, GarminConnectTooManyRequestsError
from catalog import DATASETS
from demo_data import DemoGarmin
from engine import Plan, Exporter, account_key, days, leaves, field_selected, read_methods, substitute, safe_cell


def plan(tmp_path,**kwargs):
    return Plan('2020-01-01','2020-01-02',['summary'],str(tmp_path),account_key('a@example.invalid'),delay=0,**kwargs)


def test_catalog_signatures():
    for d in DATASETS:
        f=getattr(Garmin,d.method)
        args=() if d.kind=='account' else ('2020-01-01',)* (2 if d.kind=='range' else 1)
        inspect.signature(f).bind(None,*args)


def test_multi_year_inclusive_and_leap_day():
    assert len(list(days('2019-01-01','2021-01-01')))==732
    assert list(days('2020-02-28','2020-03-01'))==['2020-02-28','2020-02-29','2020-03-01']


def test_leaf_paths_preserve_types_and_empty():
    rows=list(leaves({'a/b':[None,0,False,{},[]],'~':'=1+1'}))
    assert ('/a~1b/0',None,'null') in rows
    assert ('/a~1b/1',0,'number') in rows
    assert ('/a~1b/2',False,'boolean') in rows
    assert ('/a~1b/3',{},'object') in rows
    assert ('/~0','=1+1','string') in rows
    assert field_selected('/heartRateValues/2/1',['/heartRateValues/*/*'])
    assert not field_selected('/other/2/1',['/heartRateValues/*/*'])
    assert field_selected('/sleepDTO/stages/0',['/sleepDTO'])
    assert safe_cell(' =1+1').startswith("'")
    assert safe_cell(-8)==-8


def test_resume_and_format_change_without_redownload(tmp_path):
    a=Exporter(DemoGarmin(),plan(tmp_path,formats=['csv','json','xlsx'])); dest,report=a.run()
    assert report['stats']['downloaded']==2 and report['completed']
    source=json.loads(next(Path(dest).glob('summary-*.json')).read_text())
    assert source[0]['data']['missingExample'] is None
    with next(Path(dest).glob('summary-*.csv')).open(encoding='utf-8-sig',newline='') as f:
        rows=list(csv.DictReader(f,delimiter=';'))
    assert any(row['path']=='/missingExample' and row['type']=='null' for row in rows)
    from openpyxl import load_workbook
    workbook=load_workbook(next(Path(dest).glob('summary-*.xlsx')),read_only=True)
    assert workbook.active.max_row is None or workbook.active.max_row>2
    assert len(list(workbook.active.values))>2
    workbook.close()
    b=Exporter(DemoGarmin(),plan(tmp_path,formats=['json'])); again,r=b.run()
    assert again==dest and r['stats']['cached']==2 and r['stats']['downloaded']==0
    assert not list(Path(dest).glob('*.csv')) and not list(Path(dest).glob('*.xlsx'))


def test_account_isolation_and_filters_do_not_change_json(tmp_path):
    p=plan(tmp_path,formats=['csv','json'],selectors=['/totalSteps'])
    dest,_=Exporter(DemoGarmin(),p).run()
    with next(Path(dest).glob('summary-*.csv')).open(encoding='utf-8-sig') as f:
        assert {r['path'] for r in csv.DictReader(f,delimiter=';')}=={'/totalSteps'}
    assert 'restingHeartRate' in json.loads(next(Path(dest).glob('summary-*.json')).read_text())[0]['data']
    p.account=account_key('different@example.invalid')
    _,report=Exporter(DemoGarmin(),p).run()
    assert report['stats']['downloaded']==2


def test_cache_corruption_is_refetched(tmp_path):
    p=plan(tmp_path)
    _,r=Exporter(DemoGarmin(),p).run()
    root=Path(p.output)/('account-'+p.account)
    (root/r['records'][0]['path']).write_text('broken')
    _,r=Exporter(DemoGarmin(),p).run()
    assert r['stats']['downloaded']==1 and r['stats']['cached']==1


def test_empty_failures_rate_limit_and_cancellation(tmp_path):
    class Empty:
        def get_user_summary(self,d): return None
    _,r=Exporter(Empty(),plan(tmp_path)).run()
    assert r['stats']['empty']==2
    class Limited:
        def get_user_summary(self,d): raise GarminConnectTooManyRequestsError('sensitive message')
    _,r=Exporter(Limited(),plan(tmp_path/'limited')).run()
    assert r['fatal_error_type']=='GarminConnectTooManyRequestsError' and not r['completed']
    assert 'sensitive' not in json.dumps(r)
    stop=threading.Event();stop.set()
    _,r=Exporter(DemoGarmin(),plan(tmp_path/'cancel'),stop).run()
    assert r['cancelled'] and not r['completed']


def test_error_is_retried_next_run(tmp_path):
    class Fail:
        def get_user_summary(self,d):raise ValueError('not supported')
    _,r=Exporter(Fail(),plan(tmp_path)).run()
    assert r['stats']['errors']==2
    _,r=Exporter(DemoGarmin(),plan(tmp_path)).run()
    assert r['stats']['downloaded']==2


def test_activity_download_and_resume(tmp_path):
    class Activities:
        def get_activities_by_date(self,a,b):return [{'activityId':42}]
        def get_activity_laps(self,ident):return [{'lap':1,'distance':400}]
        def download_activity(self,ident,dl_fmt):return b'<gpx>test</gpx>'
    p=plan(tmp_path,activity_details=['get_activity_laps'],activity_files=['GPX'])
    p.datasets=['activities'];p.end=p.start
    dest,r=Exporter(Activities(),p).run()
    assert r['stats']['downloaded']==3, r
    assert len(list((Path(dest)/'activities').glob('*.gpx')))==1
    _,r=Exporter(Activities(),p).run()
    assert r['stats']['cached']==3


def test_expert_parameters_and_read_only(tmp_path):
    assert 'delete_activity' not in read_methods(Garmin)
    assert substitute({'cdate':'{date}'},'2020-01-02','x','y')=={'cdate':'2020-01-02'}
    p=plan(tmp_path,expert_method='delete_activity')
    with pytest.raises(ValueError):p.validate()
    p=plan(tmp_path,expert_method='get_heart_rate_zones_v2',expert_kwargs={'generate_defaults':True})
    with pytest.raises(ValueError):p.validate()
    p=plan(tmp_path,expert_method='get_heart_rates',expert_kwargs={'cdate':'{date}'})
    _,r=Exporter(DemoGarmin(),p).run()
    assert r['stats']['downloaded']==4


def test_wide_preserves_nested_array(tmp_path):
    p=plan(tmp_path,layout='wide');p.datasets=['heart_rate']
    dest,r=Exporter(DemoGarmin(),p).run()
    assert r['completed']
    with next(Path(dest).glob('heart_rate-*.csv')).open(encoding='utf-8-sig') as f:
        row=next(csv.DictReader(f,delimiter=';'))
    assert json.loads(row['/heartRateValues'])[2][1] is None


def test_30_activity_rows_subtype_filter_units_and_local_date(tmp_path):
    class Activities:
        def search_activities(self,start,limit,**filters):
            assert filters=={'startDate':'2020-01-01','endDate':'2020-01-01'}
            rows=[{'activityId':i,'activityName':'Testfahrt','startTimeLocal':'2020-01-01 00:30:00',
                   'startTimeGMT':'2019-12-31 23:30:00','activityType':{'typeKey':'e_bike_mountain'},
                   'distance':15670,'duration':3682,'averageSpeed':4.25,'elevationGain':216,
                   'elevationLoss':211,'avgPower':None} for i in range(30)]
            rows.append({'activityId':999,'activityType':{'typeKey':'running'}})
            return rows[start:start+20]
        def get_activity_types(self):
            return [{'typeId':1,'typeKey':'cycling'},{'typeId':2,'parentTypeId':1,'typeKey':'e_bike_mountain'}]
        def get_activities_by_date(self,*args,**kwargs):
            raise AssertionError('Use explicit unfiltered pagination')
    p=plan(tmp_path,activity_type='e_bike_mountain',formats=['csv','xlsx','json'])
    p.datasets=['activities'];p.end=p.start
    dest,r=Exporter(Activities(),p).run()
    assert r['completed'] and r['stats']['activities']==30,r
    with next(Path(dest).glob('activities-*.csv')).open(encoding='utf-8-sig',newline='') as f:
        rows=list(csv.DictReader(f,delimiter=';'))
    assert len(rows)==30
    row=rows[0]
    assert row['Datum']=='2020-01-01' and row['Startzeit']=='00:30:00'
    assert row['Distanz (km)']=='15,67' and row['Dauer (hh:mm:ss)']=='01:01:22'
    assert row['Ø Geschwindigkeit (km/h)']=='15,3' and row['Anstieg (m)']=='216'
    assert row['Ø Leistung (W)']==''
    from openpyxl import load_workbook
    book=load_workbook(next(Path(dest).glob('activities-*.xlsx')),read_only=True)
    records=list(book.active.values);book.close()
    assert len(records)==31 and records[1][5]==15.67
    assert len(json.loads(next(Path(dest).glob('activities-*.json')).read_text())[0]['data'])==30
    p.activity_type='cycling'
    _,r=Exporter(Activities(),p).run()
    assert r['stats']['cached']==1 and r['stats']['activities']==30
    p.activity_type='running'
    _,r=Exporter(Activities(),p).run()
    assert r['stats']['cached']==1 and r['stats']['activities']==1


def test_activity_pagination_repeated_page_reports_failure(tmp_path):
    class Repeated:
        def get_activities_by_date(self,*args):pass
        def search_activities(self,**kwargs):return [{'activityId':1}]
    p=plan(tmp_path);p.datasets=['activities'];p.end=p.start
    _,r=Exporter(Repeated(),p).run()
    assert r['stats']['errors']==1 and not r['completed']
    assert r['errors'][0]['error_type']=='ActivityPaginationError'


def test_activity_error_reports_status_without_private_message(tmp_path):
    from garminconnect import GarminConnectConnectionError
    class Fail:
        def get_activities_by_date(self,*args):
            raise GarminConnectConnectionError('API Error 400 - private account detail')
    p=plan(tmp_path);p.datasets=['activities'];p.end=p.start
    _,r=Exporter(Fail(),p).run()
    assert r['errors'][0]['http_status']==400
    assert 'private account detail' not in json.dumps(r)


def test_empty_activity_list_has_headers_and_no_fake_row(tmp_path):
    class Empty:
        def get_activities_by_date(self,*args):return []
    p=plan(tmp_path);p.datasets=['activities'];p.end=p.start
    dest,r=Exporter(Empty(),p).run()
    assert r['completed'] and r['stats']['activities']==0
    with next(Path(dest).glob('activities-*.csv')).open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f,delimiter=';')
        assert 'Distanz (km)' in reader.fieldnames and list(reader)==[]


def test_readable_activity_fields_filter(tmp_path):
    p=plan(tmp_path,selectors=['/*/distance','/activityName'])
    p.datasets=['activities'];p.end=p.start
    dest,r=Exporter(DemoGarmin(),p).run()
    assert r['completed']
    with next(Path(dest).glob('activities-*.csv')).open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f,delimiter=';')
        assert reader.fieldnames==['Aktivität','Distanz (km)']
        assert len(list(reader))==30


def test_connection_failure_stops_without_scanning_all_days(tmp_path):
    from garminconnect import GarminConnectConnectionError
    class Fail:
        calls=0
        def get_user_summary(self,d):
            self.calls+=1
            raise GarminConnectConnectionError('API Error 400 - private response')
    api=Fail();p=plan(tmp_path);p.end='2020-01-31'
    _,r=Exporter(api,p).run()
    assert api.calls==1 and not r['completed']
    assert r['fatal_diagnostic']['http_status']==400
    assert 'private response' not in json.dumps(r)
